"""Abas no layout "Modelo de Takt Time de VSM" (bloco de tempo disponível + tabela por produto + gráfico)."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

VERDE = PatternFill('solid', fgColor='E2EFB4')
CINZA = PatternFill('solid', fgColor='EDEDED')
BRANCO = PatternFill('solid', fgColor='FFFFFF')
lin = Side(style='thin', color='BFBFBF')
BORDA = Border(left=lin, right=lin, top=lin, bottom=lin)
F = lambda **k: Font(name='Arial', size=k.pop('size', 9), **k)
CENTRO = Alignment(horizontal='center', vertical='center', wrap_text=True)
DIREITA = Alignment(horizontal='right', vertical='center')


def _cel(ws, ref, valor=None, fill=None, fonte=None, fmt=None, alinhar=CENTRO):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    c.fill = fill or BRANCO
    c.font = fonte or F()
    c.alignment = alinhar
    if fmt:
        c.number_format = fmt
    return c


def _bloco(ws, faixa, valor, fill, fonte=None, fmt=None, alinhar=CENTRO):
    for linha in ws[faixa]:
        for c in linha:
            c.fill = fill; c.border = BORDA
    ws.merge_cells(faixa)
    return _cel(ws, faixa.split(':')[0], valor, fill, fonte, fmt, alinhar)


def aba_takt(wb, nome, titulo, rotulo_dias, dias, horas, intervalos, limpeza,
             rotulo_demanda, produtos, notas=None):
    ws = wb.create_sheet(nome)
    n = len(produtos)
    tot = L(2 + n)  # coluna TOTAIS

    ws['A1'] = titulo; ws['A1'].font = Font(name='Arial', size=18, bold=True, color='262626')
    ws['A2'] = ('Takt Time é o tempo total de produção (minutos)/demanda média (unidades). '
                'Preencha as células não sombreadas; isso preencherá as células sombreadas.')
    ws['A2'].font = F(size=8)

    # Bloco de tempo disponível
    _bloco(ws, 'A4:A4', 'TEMPO DISPONÍVEL', VERDE, F(size=8))
    _bloco(ws, 'B4:C4', 'TEMPO TOTAL', VERDE, F(size=8))
    _bloco(ws, 'D4:D5', 'INTERVALOS (min)', VERDE, F(size=8))
    _bloco(ws, 'E4:E5', 'LIMPEZA (min)', VERDE, F(size=8))
    _bloco(ws, 'F4:G5', 'MINUTOS DISPONÍVEIS', VERDE, F(size=8))
    _bloco(ws, 'A5:A5', rotulo_dias, VERDE, F(size=8))
    _bloco(ws, 'B5:C5', dias, BRANCO, F(size=11))
    _bloco(ws, 'A6:A6', 'HORAS POR DIA', VERDE, F(size=8))
    _bloco(ws, 'B6:C6', horas, BRANCO, F(size=11))
    _bloco(ws, 'D6:D6', intervalos, BRANCO, F(size=11))
    _bloco(ws, 'E6:E6', limpeza, BRANCO, F(size=11))
    _bloco(ws, 'F6:G6', '=B6*60-D6-E6', CINZA, F(size=11), '#,##0.00')

    # Tabela por produto
    _bloco(ws, 'A8:A8', 'TAKT TIME', VERDE, F(size=8))
    for j, (nome_p, _, _) in enumerate(produtos):
        _bloco(ws, f'{L(2+j)}8:{L(2+j)}8', nome_p, VERDE, F(size=8))
    _bloco(ws, f'{tot}8:{tot}8', 'TOTAIS', VERDE, F(size=8))

    rotulos = [rotulo_demanda, 'Demanda diária', 'Percentual do total de vendas',
               'Takt time (minutos por unidade)', 'Tempo de ciclo por operador (min)',
               'Quantidade de operadores']
    for i, r in enumerate(rotulos):
        _bloco(ws, f'A{9+i}:A{9+i}', r, VERDE, F(size=8))
    ws.row_dimensions[8].height = 24

    for j, (_, demanda, tc) in enumerate(produtos):
        c = L(2 + j)
        _bloco(ws, f'{c}9:{c}9', demanda, BRANCO, F(), '#,##0.00', DIREITA)
        _bloco(ws, f'{c}10:{c}10', f'=IF({c}9="","",{c}9/$B$5)', CINZA, F(), '#,##0.00', DIREITA)
        _bloco(ws, f'{c}11:{c}11', f'=IF({c}9="","",{c}9/${tot}$9)', CINZA, F(), '0%', DIREITA)
        _bloco(ws, f'{c}12:{c}12', f'=IF({c}9="","",$F$6/{c}10)', CINZA, F(), '#,##0.00', DIREITA)
        _bloco(ws, f'{c}13:{c}13', tc, BRANCO, F(), '#,##0.00', DIREITA)
        _bloco(ws, f'{c}14:{c}14', f'=IF(OR({c}13="",{c}12=""),"",{c}13/{c}12)', CINZA, F(), '#,##0.00', DIREITA)

    ult = L(1 + n)
    _bloco(ws, f'{tot}9:{tot}9', f'=SUM(B9:{ult}9)', CINZA, F(size=11), '#,##0.00', DIREITA)
    _bloco(ws, f'{tot}10:{tot}10', f'=SUM(B10:{ult}10)', CINZA, F(size=11), '#,##0.00', DIREITA)
    _bloco(ws, f'{tot}11:{tot}11', f'=SUM(B11:{ult}11)', CINZA, F(size=11), '0%', DIREITA)
    _bloco(ws, f'{tot}12:{tot}12', f'=$F$6/{tot}10', CINZA, F(size=11), '#,##0.00', DIREITA)
    _bloco(ws, f'{tot}13:{tot}13', f'=SUM(B13:{ult}13)', CINZA, F(size=11), '#,##0.00', DIREITA)
    _bloco(ws, f'{tot}14:{tot}14', f'=SUM(B14:{ult}14)', CINZA, F(size=11), '#,##0.00', DIREITA)
    ws[f'{tot}12'].comment = Comment('Takt da linha inteira: minutos disponíveis ÷ demanda diária total '
                                     '(ritmo se todos os produtos saírem da mesma linha).', 'VSM')

    # Gráfico de demanda diária
    g = BarChart(); g.title = 'DEMANDA DIÁRIA'; g.legend = None
    g.add_data(Reference(ws, min_col=2, max_col=1 + n, min_row=10), from_rows=True, titles_from_data=False)
    g.set_categories(Reference(ws, min_col=2, max_col=1 + n, min_row=8))
    g.series[0].graphicalProperties.solidFill = 'BDD7EE'
    g.series[0].graphicalProperties.line.solidFill = 'BDD7EE'
    g.y_axis.majorGridlines = g.y_axis.majorGridlines
    g.y_axis.delete = False; g.x_axis.delete = False
    g.gapWidth = 150
    g.height = 7.5; g.width = 3.2 * (n + 2)
    ws.add_chart(g, 'A16')

    if notas:
        ws[f'A32'] = notas; ws['A32'].font = F(size=8, italic=True, color='666666')

    ws.column_dimensions['A'].width = 26
    for j in range(n + 1):
        ws.column_dimensions[L(2 + j)].width = 14
    ws.sheet_view.showGridLines = False
    return ws


def montar_takt(wb):
    aba_takt(wb, 'Takt Time', 'Modelo de Takt Time de VSM', 'DIAS POR ANO', 221, 8, 30, 10,
             'Demanda anual',
             [('PRODUTO 1', 50000, 5), ('PRODUTO 2', 100000, 4), ('PRODUTO 3', 100000, 2.5),
              ('PRODUTO 4', 50000, 2), ('PRODUTO 5', 50000, 6)])
    ws = aba_takt(wb, 'Takt Thundercats', 'Takt Time — Thundercats Painéis', 'DIAS POR MÊS', 20, 8.8, 0, 0,
                  'Demanda mensal',
                  [('Lion', 170, '=9564/60'), ('Panthro', 50, None), ('Cheetara', 99, None),
                   ('Tygra', 127, None), ('Snarf', 110, None), ('WillyKit', 82, None),
                   ('WillyKat', 82, None), ('Jaga', 75, None), ('Mumm-Ra', 55, None)],
                  'Tempo de ciclo do Lion = soma dos TCs do MFV (4080 + 3812 + 822 + 850 s = 159,4 min). '
                  'A prova não dá TC das outras famílias. Detalhe por processo na aba Estado Atual.')
    ws['B13'].comment = Comment('Σ TC do MFV Atual em minutos: (4080+3812+822+850)/60', 'VSM')
