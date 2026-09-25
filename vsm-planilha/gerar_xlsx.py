from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment

AZUL = Font(name='Arial', color='0000FF')
NORMAL = Font(name='Arial')
NEG = Font(name='Arial', bold=True)
TIT = Font(name='Arial', bold=True, size=14)
CAB = Font(name='Arial', bold=True, color='FFFFFF')
FCAB = PatternFill('solid', fgColor='1A73E8')
AMARELO = PatternFill('solid', fgColor='FFF2CC')
CINZA = PatternFill('solid', fgColor='EEEEEE')
fino = Side(style='thin', color='BBBBBB')
BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)
PRIMEIRA, ULTIMA = 17, 26

def entrada(c, v):
    c.value = v; c.font = AZUL; c.fill = AMARELO; c.border = BORDA

def cabecalho(ws, linha, textos, col=1):
    for i, t in enumerate(textos):
        c = ws.cell(linha, col + i, t); c.font = CAB; c.fill = FCAB
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); c.border = BORDA

def cenario(wb, nome, titulo, param, processos, nota):
    ws = wb.create_sheet(nome)
    ws['A1'] = titulo; ws['A1'].font = TIT
    ws['A2'] = 'Células amarelas com texto azul = dados de entrada. O resto é fórmula e recalcula sozinho.'
    ws['A2'].font = Font(name='Arial', italic=True, color='666666')

    cabecalho(ws, 3, ['Parâmetro', 'Valor', 'Unidade'])
    rot = [('Demanda mensal', 'peças/mês'), ('Dias úteis', 'dias/mês'), ('Turnos da planta', 'turnos/dia'),
           ('Horas por turno', 'h'), ('Pausas por turno', 'min'), ('Estoque de produto acabado (antes da expedição)', 'peças')]
    for i, ((r, u), v) in enumerate(zip(rot, param)):
        ws.cell(4 + i, 1, r).font = NORMAL; entrada(ws.cell(4 + i, 2), v); ws.cell(4 + i, 3, u).font = NORMAL
    calc = [('Demanda diária', '=B4/B5', 'peças/dia', '0.00'),
            ('Tempo disponível por turno', '=B7*3600-B8*60', 's', '#,##0'),
            ('Tempo disponível por dia (planta)', '=B6*B11', 's/dia', '#,##0'),
            ('TAKT TIME', '=B12/B10', 's/peça', '#,##0.0')]
    for i, (r, f, u, fmt) in enumerate(calc):
        ws.cell(10 + i, 1, r).font = NEG if r == 'TAKT TIME' else NORMAL
        c = ws.cell(10 + i, 2, f); c.number_format = fmt; c.font = NEG if r == 'TAKT TIME' else NORMAL; c.border = BORDA
        ws.cell(10 + i, 3, u).font = NORMAL
    ws['A14'] = 'Takt em minutos'; ws['B14'] = '=B13/60'; ws['B14'].number_format = '0.0'; ws['C14'] = 'min/peça'
    for c in ('A14', 'B14', 'C14'): ws[c].font = NORMAL

    cabecalho(ws, 16, ['Processo', 'Operadores', 'Tempo de ciclo TC (s)', 'Lote', 'Turnos (vazio = planta)',
                       'OEE / Disponib. (vazio = 100%)', 'Setup TR (s)', 'Estoque ANTES (peças)', 'Agrega valor? (S/N)',
                       'TC efetivo = TC ÷ OEE (s)', 'Capacidade (peças/dia)', 'Carga (demanda ÷ capacidade)',
                       'Espera no estoque (dias)', 'Operadores × turnos', 'Takt (s)', 'Situação'])
    ws.row_dimensions[16].height = 48
    for i in range(PRIMEIRA, ULTIMA + 1):
        dados = processos[i - PRIMEIRA] if i - PRIMEIRA < len(processos) else [None] * 9
        for j, v in enumerate(dados):
            entrada(ws.cell(i, 1 + j), v)
        ws.cell(i, 6).number_format = '0%'
        t = f'IF(E{i}="",$B$6,E{i})'; o = f'IF(F{i}="",1,F{i})'
        f = {
            10: f'=IF(A{i}="","",C{i}/{o})',
            11: f'=IF(A{i}="","",$B$11*{t}*{o}/C{i})',
            12: f'=IF(A{i}="","",$B$10/K{i})',
            13: f'=IF(A{i}="","",H{i}/$B$10)',
            14: f'=IF(A{i}="","",B{i}*{t})',
            15: f'=IF(A{i}="","",$B$13)',
            16: f'=IF(A{i}="","",IF(L{i}>1,"NÃO ATENDE A DEMANDA",IF(J{i}>$B$13,"OK SÓ COM TURNO EXTRA","OK")))',
        }
        fmts = {10: '#,##0', 11: '0.00', 12: '0%', 13: '0.00', 14: '0', 15: '#,##0', 16: '@'}
        for col, form in f.items():
            c = ws.cell(i, col, form); c.font = NORMAL; c.number_format = fmts[col]; c.border = BORDA
        ws.cell(i, 16).number_format = 'General'

    r0 = 29
    cabecalho(ws, r0, ['Resultado', 'Valor', 'Unidade'])
    R = f'{PRIMEIRA}:{ULTIMA}'
    res = [
        ('Lead time de produção', f'=SUM(M{PRIMEIRA}:M{ULTIMA})+B9/B10', 'dias', '0.00'),
        ('Tempo de processamento (TP = Σ TC)', f'=SUM(C{PRIMEIRA}:C{ULTIMA})', 's', '#,##0'),
        ('TP em horas', f'=B{r0+2}/3600', 'h', '0.00'),
        ('Tempo de valor agregado (Σ TC com "S")', f'=SUMIF(I{PRIMEIRA}:I{ULTIMA},"S",C{PRIMEIRA}:C{ULTIMA})', 's', '#,##0'),
        ('PCE (LT em horas de trabalho)', f'=B{r0+4}/(B{r0+1}*B12+B{r0+2})', '%', '0.00%'),
        ('PCE (LT em 24 h corridas)', f'=B{r0+4}/(B{r0+1}*86400+B{r0+2})', '%', '0.00%'),
        ('Operadores teóricos (TP ÷ Takt)', f'=B{r0+2}/B13', 'operadores', '0.00'),
        ('Operadores necessários (arredondado)', f'=ROUNDUP(B{r0+7},0)', 'operadores', '0'),
        ('Operadores atuais (op. × turnos)', f'=SUM(N{PRIMEIRA}:N{ULTIMA})', 'operadores', '0'),
        ('Peças por operador por mês', f'=B4/B{r0+9}', 'peças/op/mês', '0.0'),
        ('Gargalo (maior carga)', f'=INDEX(A{PRIMEIRA}:A{ULTIMA},MATCH(MAX(L{PRIMEIRA}:L{ULTIMA}),L{PRIMEIRA}:L{ULTIMA},0))', '', 'General'),
        ('Carga do gargalo', f'=MAX(L{PRIMEIRA}:L{ULTIMA})', '%', '0%'),
    ]
    for k, (r, form, u, fmt) in enumerate(res):
        lin = r0 + 1 + k
        ws.cell(lin, 1, r).font = NORMAL
        c = ws.cell(lin, 2, form); c.number_format = fmt; c.font = NEG; c.border = BORDA; c.fill = CINZA
        ws.cell(lin, 3, u).font = NORMAL
    ws.cell(r0 + 14, 1, nota).font = Font(name='Arial', italic=True, color='666666')

    n = len(processos)
    bar = BarChart(); bar.title = 'Balanceamento: TC efetivo × Takt'; bar.y_axis.title = 'segundos'
    bar.add_data(Reference(ws, min_col=10, min_row=16, max_row=16 + n), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=17, max_row=16 + n))
    ln = LineChart(); ln.add_data(Reference(ws, min_col=15, min_row=16, max_row=16 + n), titles_from_data=True)
    ln.series[0].graphicalProperties.line.solidFill = 'D93025'
    bar += ln; bar.height = 8; bar.width = 16
    ws.add_chart(bar, 'E1')

    larg = [44, 12, 12, 8, 12, 14, 10, 12, 10, 14, 13, 14, 12, 12, 10, 24]
    for i, w in enumerate(larg):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'B17'
    return ws

def montar_cenarios(wb):
    atual = cenario(wb, 'Estado Atual', 'MFV Estado Atual — Thundercats Painéis (família Lion)',
        [170, 20, 1, 8.8, 0, 10],
        [['Puncionadeira', 1, 4080, 4, 2, 0.75, None, 35, 'S'],
         ['Dobradeira', 1, 3812, 4, 2, 0.62, None, 13, 'S'],
         ['Pré Montagem', 1, 822, None, 1, None, 477, 30, 'S'],
         ['Montagem', 1, 850, None, 1, None, 477, 25, 'S']],
        'Dados da prova. Premissa: o TC é por painel (só assim os 2 turnos da Puncionadeira e da Dobradeira se explicam).')
    atual['A20'].comment = Comment('Prova não informa pausas: considerado 0 min (8,8 h líquidas).', 'VSM')

    fut = cenario(wb, 'Estado Futuro', 'MFV Estado Futuro — proposta (edite as metas)',
        [170, 20, 1, 8.8, 0, 9],
        [['Puncionadeira (TPM + SMED)', 1, 4080, 2, 2, 0.85, None, 17, 'S'],
         ['Dobradeira (TPM + SMED) — FIFO máx. 1 lote', 1, 3812, 2, 2, 0.85, None, 4, 'S'],
         ['Célula Pré Montagem + Montagem (puxador) — supermercado antes', 1, 1672, 1, 1, None, None, 9, 'S']],
        'Metas propostas: OEE 85%, lote 2, MP 2 dias (17), FIFO 4, supermercado 1 dia (9), produto acabado 1 dia (9). Ajuste nas células amarelas.')

    ws = wb.create_sheet('Fórmulas')
    cabecalho(ws, 1, ['Indicador', 'Fórmula', 'Observação'])
    linhas = [
        ('Demanda diária', 'Demanda mensal ÷ Dias úteis', ''),
        ('Tempo disponível', 'Turnos × (Horas × 3600 − Pausas × 60)', 'Segundos por dia'),
        ('Takt time', 'Tempo disponível ÷ Demanda diária', 'Ritmo exigido pelo cliente (turnos da planta)'),
        ('TC efetivo', 'TC ÷ OEE', 'Quanto o processo realmente leva, com as perdas'),
        ('Capacidade', 'Tempo do turno × Turnos do processo × OEE ÷ TC', 'Peças por dia'),
        ('Carga', 'Demanda diária ÷ Capacidade', '> 100% = não atende; gargalo = maior carga'),
        ('Espera no estoque', 'Estoque ÷ Demanda diária', 'Dias (Lei de Little)'),
        ('Lead time', 'Σ esperas + produto acabado', 'Dias'),
        ('TP', 'Σ TC', 'Segundos'),
        ('PCE', 'Tempo VA ÷ Lead time', 'Diga qual base usou: horas de trabalho ou 24 h'),
        ('Operadores', 'Σ TC ÷ Takt', 'Arredondar para cima'),
    ]
    for i, l in enumerate(linhas):
        for j, v in enumerate(l):
            ws.cell(2 + i, 1 + j, v).font = NORMAL
    for col, w in zip('ABC', (22, 48, 50)):
        ws.column_dimensions[col].width = w


from layout_takt import montar_takt
wb = Workbook(); wb.remove(wb.active)
montar_takt(wb)
montar_cenarios(wb)
wb.save('/home/user/documentos/vsm-planilha/VSM_Thundercats.xlsx')
